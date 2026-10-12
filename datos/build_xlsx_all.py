from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment
from carreras import todas

F = "Arial"
SHORT = {"contador": "Contador Público", "administracion": "Lic. Administración", "economia": "Lic. Economía",
         "turismo": "Lic. Turismo", "profesorado": "Prof. Economía", "tec-contable": "Tec. Contable",
         "tec-comercializacion": "Tec. Comercialización", "tec-comercio-exterior": "Tec. Comercio Exterior",
         "tec-periodismo": "Tec. Periodismo Digital", "tec-turismo": "Tec. Turismo"}
H = ["Ciclo", "Orientación / Grupo", "Año", "Cuat.", "Área", "Cód.", "Denominación", "Grupo", "Hs/s", "Hs Bim",
     "Hs Cuat", "Corr. 1", "Corr. 2", "Corr. 3", "Corr. 4", "Corr. 5", "Estado", "¿Puedo cursar?", "Correlativas pendientes", "¿Puedo rendir el final?", "Correlativas a revisar"]
REGLA = "Resumen!$B$15"
num = lambda x: int(x) if isinstance(x, str) and x.isdigit() else (x if x not in ("", None) else None)
fill = lambda h: PatternFill("solid", fgColor=h, bgColor=h)
thin = Side(style="thin", color="D0D7E2")

wb = Workbook()
res = wb.active; res.title = "Resumen"
res["A1"] = "UNMdP Map · Ciencias Económicas y Sociales"; res["A1"].font = Font(name=F, bold=True, size=14)
for j, h in enumerate(["Carrera", "Plan", "Materias", "Aprobadas", "Cursadas aprobadas", "Puedo cursar", "Bloqueadas"], 1):
    c = res.cell(3, j, h); c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = fill("1F3A5F")

for k, car in enumerate(todas()):
    ms = car["materias"]
    assert max(len(m["corr"]) for m in ms) <= 5
    ws = wb.create_sheet(SHORT[car["id"]])
    ws.append(H)
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = fill("1F3A5F")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    last = len(ms) + 1
    codes, st = f"$F$2:$F${last}", f"$Q$2:$Q${last}"
    for i, m in enumerate(ms, 2):
        anio = "Optativa" if m["anio"].startswith("O") else num(m["anio"])
        corr = [num(c) for c in m["corr"]] + [None] * (5 - len(m["corr"]))
        ws.append([m["ciclo"], m["orientacion"], anio, num(m["cuat"]), m["area"], num(m["cod"]), m["nombre"],
                   m["grupo"], m["hs_sem"], m["hs_bim"], m["hs_cuat"], *corr, "No cursada"])
        cols = "LMNOP"
        fin = lambda c: f'COUNTIFS({codes},{c}{i},{st},"Aprobada")>0'
        cur = lambda c: f'OR({c}{i}="",COUNTIF({codes},{c}{i})=0,{fin(c)},AND({REGLA}="La cursada aprobada",COUNTIFS({codes},{c}{i},{st},"Cursada aprobada")>0))'
        rin = lambda c: f'OR({c}{i}="",COUNTIF({codes},{c}{i})=0,{fin(c)})'
        ws[f"R{i}"] = f'=IF(Q{i}="Aprobada","Ya aprobada",IF(Q{i}="Cursada aprobada","Ya cursada",IF(AND({",".join(cur(c) for c in cols)}),"Sí","No")))'
        ws[f"S{i}"] = "=TRIM(" + "&".join(f'IF(NOT({cur(c)}),{c}{i}&" ","")' for c in cols) + ")"
        ws[f"T{i}"] = f'=IF(Q{i}="Cursada aprobada",IF(AND({",".join(rin(c) for c in cols)}),"Sí","No"),"")'
        # control: una materia aprobada exige correlativas aprobadas; una cursada exige que cumplan la regla para cursar
        noFin = "&".join(f'IF(NOT({rin(c)}),{c}{i}&" ","")' for c in cols)
        noCur = "&".join(f'IF(NOT({cur(c)}),{c}{i}&" ","")' for c in cols)
        ws[f"U{i}"] = f'=IF(Q{i}="Aprobada",TRIM({noFin}),IF(Q{i}="Cursada aprobada",TRIM({noCur}),""))'
        for c in ws[i]:
            c.font = Font(name=F); c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="center", horizontal="left" if c.column in (1, 2, 7, 19) else "center")
        ws[f"Q{i}"].font = Font(name=F, bold=True, color="0000FF")
        missing = [c for c in m["corr"] if c not in {x["cod"] for x in ms}]
        if missing:
            ws[f"G{i}"].comment = Comment(f"Correlativa {', '.join(missing)} no figura en este plan; no se tiene en cuenta.", "Plan")
    dv = DataValidation(type="list", formula1='"No cursada,Cursada aprobada,Aprobada"', allow_blank=False, showErrorMessage=True,
                        errorTitle="Valor no válido", error="Elegí No cursada, Cursada aprobada o Aprobada.")
    ws.add_data_validation(dv); dv.add(f"Q2:Q{last}")
    ws.conditional_formatting.add(f"Q2:Q{last}", FormulaRule(formula=['$Q2="Aprobada"'], fill=fill("C8E6C9")))
    ws.conditional_formatting.add(f"Q2:Q{last}", FormulaRule(formula=['$Q2="Cursada aprobada"'], fill=fill("D4EEF4")))
    ws.conditional_formatting.add(f"U2:U{last}", FormulaRule(formula=['LEN($U2)>0'], fill=fill("FFCDD2"), font=Font(bold=True, color="B71C1C")))
    ws.conditional_formatting.add(f"T2:T{last}", FormulaRule(formula=['$T2="Sí"'], fill=fill("C8E6C9"), font=Font(bold=True, color="1B5E20")))
    ws.conditional_formatting.add(f"T2:T{last}", FormulaRule(formula=['$T2="No"'], fill=fill("FFCDD2"), font=Font(color="B71C1C")))
    ws.conditional_formatting.add(f"R2:R{last}", FormulaRule(formula=['$R2="Sí"'], fill=fill("C8E6C9"), font=Font(bold=True, color="1B5E20")))
    ws.conditional_formatting.add(f"R2:R{last}", FormulaRule(formula=['$R2="No"'], fill=fill("FFCDD2"), font=Font(color="B71C1C")))
    ws.conditional_formatting.add(f"R2:R{last}", FormulaRule(formula=['OR($R2="Ya aprobada",$R2="Ya cursada")'], fill=fill("ECEFF1"), font=Font(color="607D8B")))
    for col, w in zip("ABCDEFGHIJKLMNOPQRSTU", [24, 30, 9, 7, 7, 8, 50, 7, 6, 8, 8, 8, 8, 8, 8, 8, 17, 15, 22, 15, 18]):
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "H2"; ws.auto_filter.ref = f"A1:U{last}"
    ws["Q1"].comment = Comment("Desplegable: No cursada / Cursada aprobada / Aprobada. Es la única columna para editar.", "Plan")
    ws["U1"].comment = Comment("Avisa en rojo si la materia figura Aprobada con alguna correlativa sin aprobar, o Cursada aprobada con correlativas que no cumplen la regla para cursar. Corregí el estado de la materia o de esas correlativas.", "Plan")
    ws["T1"].comment = Comment("Solo para materias con la cursada aprobada: Sí cuando todas las correlativas están Aprobada.", "Plan")
    n = len(ms) + 1
    nn = f"'{ws.title}'"
    r = 4 + k
    vals = [car["nombre"], car["plan"], f"=COUNTA({nn}!F2:F{n})", f'=COUNTIF({nn}!Q2:Q{n},"Aprobada")',
            f'=COUNTIF({nn}!Q2:Q{n},"Cursada aprobada")', f'=COUNTIF({nn}!R2:R{n},"Sí")', f'=COUNTIF({nn}!R2:R{n},"No")']
    for j, v in enumerate(vals, 1):
        res.cell(r, j, v).font = Font(name=F)
    res.cell(r, 1).hyperlink = car["url"]

res["A15"] = "Para cursar, las correlativas tienen que tener:"; res["A15"].font = Font(name=F, bold=True)
res["B15"] = "La cursada aprobada"; res["B15"].font = Font(name=F, bold=True, color="0000FF"); res["B15"].fill = fill("FFF59D")
dvr = DataValidation(type="list", formula1='"La cursada aprobada,El final aprobado"', allow_blank=False)
res.add_data_validation(dvr); dvr.add("B15")
res["B15"].comment = Comment("Los planes no indican si para cursar alcanza con la cursada aprobada de las correlativas o se pide el final. Elegí la regla que corresponda; para rendir el final siempre se piden las correlativas aprobadas.", "Plan")
notes = ["Cómo usar: en la hoja de tu carrera cambiá la columna Estado (texto azul): No cursada, Cursada aprobada o Aprobada.",
         "¿Puedo cursar? muestra Sí cuando las correlativas cumplen la regla elegida arriba (celda amarilla).",
         "¿Puedo rendir el final? aparece en las materias con la cursada aprobada: Sí cuando todas sus correlativas están Aprobada.",
         "Una materia nunca puede estar Aprobada si alguna correlativa no lo está. La columna Correlativas a revisar marca en rojo esos casos.",
         "La cursada aprobada tiene un año de vigencia (OCA 357/2025).",
         "Los códigos de correlativas que no figuran en el plan no se tienen en cuenta (tienen un comentario).",
         "Fuente: planes de estudio publicados en eco.mdp.edu.ar (clic en el nombre de la carrera). Verificá con la facultad."]
for i, t in enumerate(notes, 17):
    res[f"A{i}"] = t; res[f"A{i}"].font = Font(name=F)
for col, w in zip("ABCDEFG", [44, 36, 10, 11, 18, 13, 12]):
    res.column_dimensions[col].width = w
wb.save("unmdp-map_fceys.xlsx")
print("ok")
